"""Generated Excel template transaction-type dropdown. Template only."""

from __future__ import annotations

from datetime import date, datetime

import openpyxl


EXPECTED_HEADERS = [
    "Date",
    "Borrower Name",
    "Due Day",
    "Transaction Type",
    "Amount",
    "Payment Mode",
    "Notes",
]

ALLOWED_TXN_TYPES = ["Interest Received", "Principal Received"]


def _list_values(formula1) -> list[str]:
    text = str(formula1 or "").strip()
    if text.startswith('"') and text.endswith('"'):
        text = text[1:-1]
    return [part.strip() for part in text.split(",") if part.strip()]


def test_generated_template_transaction_types(initialized_temp_db, tmp_path):
    import database as db
    from src.services.import_template_generator import ImportTemplateGenerator
    from src.services.loan_matcher import LoanMatcher

    db.add_borrower("Template Borrower", "", "")
    out = tmp_path / "txn_types_template.xlsx"
    matcher = LoanMatcher(initialized_temp_db)
    ok = ImportTemplateGenerator(initialized_temp_db, str(out)).generate(matcher)
    assert ok is True

    wb = openpyxl.load_workbook(out)
    assert "Daily Transactions" in wb.sheetnames
    ws = wb["Daily Transactions"]
    assert [c.value for c in ws[1]] == EXPECTED_HEADERS

    sample = ws["A2"]
    assert isinstance(sample.value, (date, datetime))
    assert "YY" in (sample.number_format or "").upper()

    date_validations = [
        dv for dv in ws.data_validations.dataValidation if dv.type == "date"
    ]
    assert date_validations
    assert any("A2" in str(dv.sqref) for dv in date_validations)

    list_validations = [
        dv
        for dv in ws.data_validations.dataValidation
        if dv.type == "list" and dv.sqref and "D2" in str(dv.sqref)
    ]
    assert list_validations, "Transaction Type column should have a list dropdown"
    values = _list_values(list_validations[0].formula1)
    assert values == ALLOWED_TXN_TYPES
    assert "Loan Given" not in values
    assert "Adjustment" not in values
    joined = ",".join(values)
    assert "Loan Given" not in joined
    assert "Adjustment" not in joined
    assert "Loan Given" not in (list_validations[0].error or "")
    assert "Adjustment" not in (list_validations[0].error or "")
