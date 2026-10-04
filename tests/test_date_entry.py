"""Date picker and Excel Date-column usability tests. UI/template only."""

from __future__ import annotations

from datetime import date, datetime

import customtkinter as ctk
import openpyxl

from src.ui.date_picker import (
    DatePickerDialog,
    format_iso_date,
    parse_entry_date,
    set_entry_iso,
)


def test_format_iso_date():
    assert format_iso_date(date(2026, 10, 4)) == "2026-10-04"


def test_parse_entry_date_valid_and_invalid():
    assert parse_entry_date("2026-04-15") == date(2026, 4, 15)
    assert parse_entry_date("not-a-date", fallback=date(2020, 1, 1)) == date(2020, 1, 1)
    assert parse_entry_date("2026-13-40", fallback=date(2021, 2, 2)) == date(2021, 2, 2)


def test_set_entry_iso_and_manual_typing():
    ctk.set_appearance_mode("dark")
    root = ctk.CTk()
    root.withdraw()
    entry = ctk.CTkEntry(root)
    entry.insert(0, "2020-01-01")
    set_entry_iso(entry, "2026-04-15")
    assert entry.get() == "2026-04-15"
    entry.delete(0, "end")
    entry.insert(0, "2024-12-31")
    assert entry.get() == "2024-12-31"
    root.destroy()


def test_calendar_select_writes_iso_without_changing_invalid_until_pick():
    ctk.set_appearance_mode("dark")
    root = ctk.CTk()
    root.withdraw()
    entry = ctk.CTkEntry(root)
    entry.insert(0, "typed-manually")
    dlg = DatePickerDialog(
        root,
        initial=date(2026, 3, 1),
        on_select=lambda iso: set_entry_iso(entry, iso),
    )
    root.update()
    assert entry.get() == "typed-manually"
    dlg._select_day(9)
    root.update()
    assert entry.get() == "2026-03-09"
    root.destroy()


def test_calendar_cancel_leaves_field_unchanged():
    ctk.set_appearance_mode("dark")
    root = ctk.CTk()
    root.withdraw()
    entry = ctk.CTkEntry(root)
    entry.insert(0, "2025-06-01")
    dlg = DatePickerDialog(
        root,
        initial=date(2026, 1, 1),
        on_select=lambda iso: set_entry_iso(entry, iso),
    )
    root.update()
    dlg._cancel()
    root.update()
    assert entry.get() == "2025-06-01"
    root.destroy()


def test_excel_template_date_column_is_native_date(initialized_temp_db, tmp_path):
    import database as db
    from src.services.import_template_generator import ImportTemplateGenerator
    from src.services.loan_matcher import LoanMatcher

    db.add_borrower("Template Borrower", "", "")
    out = tmp_path / "template.xlsx"
    matcher = LoanMatcher(initialized_temp_db)
    ok = ImportTemplateGenerator(initialized_temp_db, str(out)).generate(matcher)
    assert ok is True

    wb = openpyxl.load_workbook(out)
    ws = wb["Daily Transactions"]
    assert [c.value for c in ws[1]] == [
        "Date",
        "Borrower Name",
        "Due Day",
        "Transaction Type",
        "Amount",
        "Payment Mode",
        "Notes",
    ]
    sample = ws["A2"]
    assert isinstance(sample.value, (date, datetime))
    assert "YY" in (sample.number_format or "").upper()
    assert ws["A10"].number_format == sample.number_format

    date_validations = [
        dv for dv in ws.data_validations.dataValidation if dv.type == "date"
    ]
    assert date_validations, "Date column should have Excel date validation"
    ranges = " ".join(str(dv.sqref) for dv in date_validations)
    assert "A2" in ranges
