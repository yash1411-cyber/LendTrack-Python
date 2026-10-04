"""
ui_reports.py - Borrower-wise summary report + Excel export.
"""

import customtkinter as ctk
from tkinter import filedialog
from datetime import date
import database as db
from src.ui.theme import (
    BG, CARD, CARD2, TEXT, MUTED, BORDER, ACCENT, GREEN, RED,
    font_body, font_overline, font_body_bold, CONTENT_PAD,
    S8, S12, S16, S32, RADIUS_CARD, ROW_PY,
    fmt_inr, success_button, secondary_button,
    table_header, table_cell, bind_row_hover, ellipsize, notify,
)


class ReportsFrame(ctk.CTkFrame):
    def __init__(self, parent, **kwargs):
        super().__init__(parent, fg_color=BG, **kwargs)
        self._report_data = []
        self._build_ui()

    def refresh(self):
        self._load()

    def _build_ui(self):
        self.grid_columnconfigure(0, weight=1)
        self.grid_rowconfigure(2, weight=1)

        hdr = ctk.CTkFrame(self, fg_color="transparent")
        hdr.grid(row=0, column=0, sticky="ew", padx=CONTENT_PAD, pady=(S16, S8))
        success_button(hdr, "Export to Excel", self._export_excel, width=160).pack(
            side="right"
        )
        secondary_button(hdr, "Refresh", self._load, width=100).pack(
            side="right", padx=(0, 10)
        )

        self._summary_frame = ctk.CTkFrame(
            self, fg_color=CARD, corner_radius=RADIUS_CARD,
            border_width=1, border_color=BORDER,
        )
        self._summary_frame.grid(row=1, column=0, sticky="ew", padx=CONTENT_PAD, pady=(0, 4))
        self._summary_frame.grid_columnconfigure((0, 1, 2, 3, 4), weight=1)
        self._sum_labels = {}
        for ci, (key, lbl) in enumerate([
            ("total_principal", "Total Outstanding"),
            ("total_given", "Total Loans Given"),
            ("total_returned", "Total Returned"),
            ("total_interest_recv", "Interest Received"),
            ("total_pending", "Pending Interest"),
        ]):
            f = ctk.CTkFrame(self._summary_frame, fg_color="transparent")
            f.grid(row=0, column=ci, padx=S16, pady=S12)
            ctk.CTkLabel(
                f, text=lbl, font=font_overline(), text_color=MUTED, wraplength=140
            ).pack()
            vl = ctk.CTkLabel(
                f, text=fmt_inr(0), font=font_body_bold(), text_color=TEXT
            )
            vl.pack()
            self._sum_labels[key] = vl

        content = ctk.CTkFrame(self, fg_color="transparent")
        content.grid(row=2, column=0, sticky="nsew", padx=CONTENT_PAD, pady=(S8, S16))
        content.grid_columnconfigure(0, weight=1)
        content.grid_rowconfigure(0, weight=1)

        cols = [
            "ID", "Name", "Phone",
            "Principal Given", "Returned",
            "Outstanding", "Exp. Interest/Mo",
            "Interest Recv'd", "Pending Interest",
        ]
        weights = [1, 2, 1, 2, 2, 2, 2, 2, 2]

        tbl = ctk.CTkScrollableFrame(
            content, fg_color=CARD, corner_radius=RADIUS_CARD,
            border_width=1, border_color=BORDER,
        )
        tbl.grid(row=0, column=0, sticky="nsew")
        for ci, w in enumerate(weights):
            tbl.grid_columnconfigure(ci, weight=w)
        self._tbl = tbl

        for ci, col in enumerate(cols):
            table_header(tbl, col).grid(
                row=0, column=ci, padx=S12, pady=10, sticky="w"
            )
        self._row_widgets = []
        self._load()

    def _load(self):
        for w in self._row_widgets:
            w.destroy()
        self._row_widgets.clear()

        self._report_data = db.get_report_data()

        tp = sum(r["outstanding_principal"] for r in self._report_data)
        tg = sum(r["total_principal_given"] for r in self._report_data)
        tr = sum(r["principal_returned"] for r in self._report_data)
        tir = sum(r["interest_received"] for r in self._report_data)
        tpd = sum(r["pending_interest"] for r in self._report_data)
        for key, val in [
            ("total_principal", tp), ("total_given", tg),
            ("total_returned", tr), ("total_interest_recv", tir),
            ("total_pending", tpd),
        ]:
            c = RED if key == "total_pending" and val > 0 else GREEN if key in ("total_interest_recv",) else TEXT
            self._sum_labels[key].configure(text=fmt_inr(val), text_color=c)

        if not self._report_data:
            lbl = ctk.CTkLabel(
                self._tbl, text="No data yet.", text_color=MUTED, font=font_body()
            )
            lbl.grid(row=1, column=0, columnspan=9, pady=S32)
            self._row_widgets.append(lbl)
            return

        for ri, r in enumerate(self._report_data, start=1):
            bg = CARD if ri % 2 == 0 else CARD2
            rf = ctk.CTkFrame(self._tbl, fg_color=bg, corner_radius=0)
            rf.grid(row=ri, column=0, columnspan=9, sticky="ew")
            for ci in range(9):
                rf.grid_columnconfigure(ci, weight=[1, 2, 1, 2, 2, 2, 2, 2, 2][ci])
            bind_row_hover(rf, bg)
            self._row_widgets.append(rf)

            pend_col = RED if r["pending_interest"] > 0 else TEXT
            vals = [
                r["borrower_id"], ellipsize(r["name"], 20), ellipsize(r["phone"] or "—", 14),
                fmt_inr(r["total_principal_given"]),
                fmt_inr(r["principal_returned"]),
                fmt_inr(r["outstanding_principal"]),
                fmt_inr(r["expected_monthly_interest"]),
                fmt_inr(r["interest_received"]),
                fmt_inr(r["pending_interest"]),
            ]
            colors = [MUTED, TEXT, MUTED, TEXT, TEXT, ACCENT, GREEN, GREEN, pend_col]
            for ci, (v, col) in enumerate(zip(vals, colors)):
                table_cell(rf, v, col, row=0, column=ci, padx=S12, pady=ROW_PY, sticky="w")

    def _export_excel(self):
        if not self._report_data:
            notify(self, "Nothing to export.", "There is no report data to export.", tone="info")
            return
        try:
            import openpyxl
            from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        except ImportError:
            notify(
                self,
                "Excel export is unavailable.",
                "openpyxl is required for Excel export.\n"
                "Run:  pip install openpyxl",
                tone="error",
            )
            return

        path = filedialog.asksaveasfilename(
            defaultextension=".xlsx",
            filetypes=[("Excel files", "*.xlsx")],
            initialfile=f"LendTrack_Report_{date.today()}.xlsx"
        )
        if not path:
            return

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Borrower Summary"

        hdr_fill = PatternFill("solid", fgColor="1A2030")
        hdr_font = Font(bold=True, color="4F8EF7", size=11)
        border = Border(bottom=Side(style="thin", color="2C3347"))
        center = Alignment(horizontal="center")

        headers = [
            "ID", "Name", "Phone",
            "Principal Given", "Returned",
            "Outstanding", "Exp Interest/Mo",
            "Interest Received", "Pending Interest",
        ]
        ws.append(headers)
        for cell in ws[1]:
            cell.font = hdr_font
            cell.fill = hdr_fill
            cell.alignment = center
            cell.border = border

        for r in self._report_data:
            ws.append([
                r["borrower_id"], r["name"], r["phone"] or "",
                r["total_principal_given"], r["principal_returned"],
                r["outstanding_principal"], r["expected_monthly_interest"],
                r["interest_received"], r["pending_interest"],
            ])

        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = 18

        wb.save(path)
        notify(
            self,
            "Report exported.",
            f"The report was saved to:\n{path}",
            tone="success",
        )
